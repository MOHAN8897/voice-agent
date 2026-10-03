"""Intelligent CSV/Excel contact parsing, column mapping, phone normalization, and variable resolution engine."""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import uuid
from datetime import datetime, timezone
from typing import Any

logger = logging.getLogger(__name__)

# Supported canonical Voxly fields
CANONICAL_FIELDS: dict[str, dict[str, Any]] = {
    "phone": {"label": "Phone Number", "category": "required", "description": "Primary phone in E.164 format"},
    "first_name": {"label": "First Name", "category": "recommended", "description": "Contact's given name"},
    "last_name": {"label": "Last Name", "category": "recommended", "description": "Contact's family name / surname"},
    "full_name": {"label": "Full Name", "category": "recommended", "description": "Contact's full name"},
    "email": {"label": "Email", "category": "recommended", "description": "Contact's email address"},
    "company": {"label": "Company", "category": "recommended", "description": "Company or organization name"},
    "job_title": {"label": "Job Title", "category": "recommended", "description": "Job title or designation"},
    "country": {"label": "Country", "category": "recommended", "description": "Country of residence"},
    "city": {"label": "City", "category": "recommended", "description": "City or town"},
    "state": {"label": "State", "category": "recommended", "description": "State / province / region"},
    "timezone": {"label": "Timezone", "category": "recommended", "description": "Timezone e.g. America/New_York"},
    "notes": {"label": "Notes", "category": "recommended", "description": "Call notes or remarks"},
    "address": {"label": "Address", "category": "optional", "description": "Street address"},
    "website": {"label": "Website", "category": "optional", "description": "Website URL"},
    "customer_id": {"label": "Customer ID", "category": "optional", "description": "Internal CRM or customer ID"},
    "lead_id": {"label": "Lead ID", "category": "optional", "description": "Lead identifier"},
    "language": {"label": "Language", "category": "optional", "description": "Preferred language"},
    "industry": {"label": "Industry", "category": "optional", "description": "Business industry"},
    "source": {"label": "Source", "category": "optional", "description": "Lead source or channel"},
    "tags": {"label": "Tags", "category": "optional", "description": "Comma-separated tags"},
}

# Aliases dictionary
FIELD_ALIASES: dict[str, list[str]] = {
    "phone": [
        "phone", "phone_number", "phone number", "phonenumber", "phone no", "phone_no", "phoneno",
        "mobile", "mobile_number", "mobile number", "mobilenumber", "mobile no", "mobile_no", "mobileno",
        "telephone", "telephone_number", "telephone number", "telephonenumber", "tel", "tel_no", "tel no",
        "contact", "contact_number", "contact number", "contactnumber", "contact_no", "contact no", "contactno",
        "cell", "cell_number", "cell number", "cellphone", "cell phone", "cellular",
        "whatsapp", "whatsapp_number", "whatsapp number", "wa_number",
    ],
    "full_name": [
        "name", "full_name", "full name", "fullname", "customer_name", "customer name", "customername",
        "contact_name", "contact name", "contactname", "lead_name", "lead name", "leadname",
        "client_name", "client name", "clientname", "person_name", "person name",
    ],
    "first_name": [
        "first_name", "first name", "firstname", "fname", "given_name", "given name", "givenname",
        "forename",
    ],
    "last_name": [
        "last_name", "last name", "lastname", "lname", "surname", "family_name", "family name",
        "familyname",
    ],
    "email": [
        "email", "email_address", "email address", "emailaddress", "email_id", "email id", "emailid",
        "mail", "mail_id", "mail id", "e_mail", "e mail", "contact_email", "contact email",
    ],
    "company": [
        "company", "company_name", "company name", "companyname", "organization", "organisation",
        "business", "business_name", "business name", "businessname", "employer", "org", "firm",
        "account_name", "account name",
    ],
    "job_title": [
        "job_title", "job title", "jobtitle", "title", "designation", "role", "position", "occupation",
    ],
    "city": ["city", "town", "municipality"],
    "state": ["state", "province", "region", "territory"],
    "country": ["country", "country_name", "country name", "country_code", "country code", "nation"],
    "timezone": ["timezone", "time_zone", "time zone", "tz"],
    "notes": ["notes", "note", "comments", "comment", "remarks", "remark", "description", "details", "memo"],
    "customer_id": ["customer_id", "customer id", "customerid", "cust_id", "cust id", "client_id"],
    "lead_id": ["lead_id", "lead id", "leadid"],
    "address": ["address", "street", "street_address", "street address", "location"],
    "website": ["website", "web", "url", "site"],
    "language": ["language", "lang", "locale"],
    "industry": ["industry", "sector", "vertical"],
    "source": ["source", "lead_source", "lead source", "channel"],
    "tags": ["tags", "tag", "labels", "label"],
}

# Inverted alias map for fast lookup: normalized alias -> field
NORMALIZED_ALIAS_MAP: dict[str, str] = {}
for canonical, aliases in FIELD_ALIASES.items():
    for alias in aliases:
        norm = re.sub(r"[^a-z0-9]", "", alias.lower())
        if norm and norm not in NORMALIZED_ALIAS_MAP:
            NORMALIZED_ALIAS_MAP[norm] = canonical

# Country prefix codes
DEFAULT_COUNTRY_CODES: dict[str, str] = {
    "US": "1",
    "CA": "1",
    "IN": "91",
    "GB": "44",
    "UK": "44",
    "AU": "61",
    "AE": "971",
    "SG": "65",
    "DE": "49",
    "FR": "33",
    "ZA": "27",
    "NZ": "64",
}


def normalize_header_text(header: str) -> str:
    """Normalize header by trimming, lowercasing, and removing punctuation."""
    if not header:
        return ""
    # Replace separators with single space
    cleaned = re.sub(r"[-_/]+", " ", str(header).strip().lower())
    # Remove remaining punctuation except alphanumeric and space
    cleaned = re.sub(r"[^\w\s]", "", cleaned)
    return " ".join(cleaned.split())


def _is_probable_email(val: str) -> bool:
    if not val or "@" not in val or "." not in val:
        return False
    return bool(re.match(r"^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$", val.strip()))


def _is_probable_phone(val: str) -> bool:
    if not val:
        return False
    digits = re.sub(r"[^\d]", "", val)
    # Most valid phone numbers worldwide have between 7 and 15 digits
    return 7 <= len(digits) <= 15 and (val.strip().startswith("+") or len(digits) >= 10)


def _is_probable_name(val: str) -> bool:
    if not val or len(val) < 2 or len(val) > 70:
        return False
    # Names typically don't have @ or leading digits
    if "@" in val or val[0].isdigit():
        return False
    # Must contain letters
    return bool(re.search(r"[a-zA-Z\u0C00-\u0C7F]", val))


def parse_file_to_rows(
    file_bytes: bytes,
    file_name: str,
    max_preview_rows: int = 100,
) -> tuple[list[str], list[dict[str, Any]], int, list[str]]:
    """Parse CSV or Excel file.
    Returns: (headers, preview_rows, total_row_count, warning_messages)
    """
    file_name_lower = (file_name or "").lower()
    headers: list[str] = []
    rows: list[dict[str, Any]] = []
    warnings: list[str] = []

    if file_name_lower.endswith(".xlsx") or file_name_lower.endswith(".xls"):
        # Excel parsing
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
            sheet = wb.active
            if sheet is None:
                raise ValueError("Spreadsheet contains no active sheet")

            all_iter = sheet.iter_rows(values_only=True)
            try:
                first_row = next(all_iter)
            except StopIteration:
                return [], [], 0, ["The uploaded spreadsheet is empty."]

            raw_headers = [str(col).strip() if col is not None else "" for col in first_row]
            # Clean and deduplicate headers
            seen_headers: set[str] = set()
            col_indices: list[int] = []

            for idx, h in enumerate(raw_headers):
                if not h:
                    h = f"Column_{idx + 1}"
                original_h = h
                count = 1
                while h.lower() in seen_headers:
                    warnings.append(f"Duplicate header detected: '{original_h}'. Renamed to '{original_h}_{count}'.")
                    h = f"{original_h}_{count}"
                    count += 1
                seen_headers.add(h.lower())
                headers.append(h)
                col_indices.append(idx)

            total_count = 0
            for row_vals in all_iter:
                # check if completely empty row
                if not any(row_vals):
                    continue
                total_count += 1
                if len(rows) < max_preview_rows:
                    row_dict: dict[str, Any] = {}
                    for col_idx, header_name in zip(col_indices, headers):
                        val = row_vals[col_idx] if col_idx < len(row_vals) else None
                        row_dict[header_name] = str(val).strip() if val is not None else ""
                    rows.append(row_dict)

            wb.close()
            return headers, rows, total_count, warnings

        except Exception as e:
            logger.error("Excel parse failed: %s", e)
            raise ValueError(f"Could not parse Excel spreadsheet: {e}")

    else:
        # CSV parsing with delimiter detection and encoding fallbacks
        text_content = ""
        for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                text_content = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if not text_content:
            raise ValueError("Unsupported file encoding. Please ensure the CSV is encoded in UTF-8.")

        # Detect delimiter
        sample = text_content[:4096]
        delimiter = ","
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
            delimiter = dialect.delimiter
        except Exception:
            if "\t" in sample:
                delimiter = "\t"
            elif ";" in sample:
                delimiter = ";"

        reader = csv.reader(io.StringIO(text_content), delimiter=delimiter)
        try:
            first_row = next(reader)
        except StopIteration:
            return [], [], 0, ["The uploaded CSV file is empty."]

        raw_headers = [str(col).strip() for col in first_row]
        seen_headers = set()
        for idx, h in enumerate(raw_headers):
            if not h:
                h = f"Column_{idx + 1}"
            original_h = h
            count = 1
            while h.lower() in seen_headers:
                warnings.append(f"Duplicate header detected: '{original_h}'. Renamed to '{original_h}_{count}'.")
                h = f"{original_h}_{count}"
                count += 1
            seen_headers.add(h.lower())
            headers.append(h)

        total_count = 0
        for row_vals in reader:
            if not any(row_vals):
                continue
            total_count += 1
            if len(rows) < max_preview_rows:
                row_dict = {}
                for idx, h in enumerate(headers):
                    val = row_vals[idx] if idx < len(row_vals) else ""
                    row_dict[h] = str(val).strip()
                rows.append(row_dict)

        return headers, rows, total_count, warnings


def detect_column_mappings(
    headers: list[str],
    sample_rows: list[dict[str, Any]],
    saved_template_mapping: dict[str, str] | None = None,
) -> dict[str, dict[str, Any]]:
    """Detect mappings with confidence scores for each uploaded header.
    Returns: header -> {
        "field": canonical_field or custom_field_name,
        "is_custom": bool,
        "confidence": int (0-100),
        "confidence_level": "HIGH" | "MEDIUM" | "LOW",
        "sample_values": list[str],
        "reason": str
    }
    """
    results: dict[str, dict[str, Any]] = {}
    assigned_canonicals: set[str] = set()

    # Pre-extract sample values per header (3-5 non-empty values)
    samples_per_header: dict[str, list[str]] = {}
    for h in headers:
        samples = []
        for r in sample_rows:
            v = str(r.get(h) or "").strip()
            if v and v not in samples:
                samples.append(v)
            if len(samples) >= 5:
                break
        samples_per_header[h] = samples

    for h in headers:
        samples = samples_per_header.get(h, [])
        norm_h = normalize_header_text(h)
        compressed_h = re.sub(r"[^a-z0-9]", "", norm_h)

        # 1. Check saved template mapping first
        if saved_template_mapping and (h in saved_template_mapping or norm_h in saved_template_mapping):
            target = saved_template_mapping.get(h) or saved_template_mapping.get(norm_h)
            is_custom = target not in CANONICAL_FIELDS and target != "ignore"
            results[h] = {
                "field": target,
                "is_custom": is_custom,
                "confidence": 98,
                "confidence_level": "HIGH",
                "sample_values": samples,
                "reason": "Matched saved import template",
            }
            if not is_custom and target != "ignore":
                assigned_canonicals.add(target)
            continue

        # 2. Check ambiguous column names like "Contact"
        if compressed_h in ("contact", "contactdetails", "contacts"):
            # inspect sample values
            phone_votes = sum(1 for s in samples if _is_probable_phone(s))
            email_votes = sum(1 for s in samples if _is_probable_email(s))
            name_votes = sum(1 for s in samples if _is_probable_name(s))

            if phone_votes > email_votes and phone_votes > name_votes:
                results[h] = {
                    "field": "phone",
                    "is_custom": False,
                    "confidence": 88,
                    "confidence_level": "MEDIUM",
                    "sample_values": samples,
                    "reason": "Header 'Contact' contained phone number values",
                }
                assigned_canonicals.add("phone")
                continue
            elif email_votes > phone_votes and email_votes > name_votes:
                results[h] = {
                    "field": "email",
                    "is_custom": False,
                    "confidence": 88,
                    "confidence_level": "MEDIUM",
                    "sample_values": samples,
                    "reason": "Header 'Contact' contained email address values",
                }
                assigned_canonicals.add("email")
                continue
            elif name_votes > phone_votes and name_votes > email_votes:
                results[h] = {
                    "field": "full_name",
                    "is_custom": False,
                    "confidence": 85,
                    "confidence_level": "MEDIUM",
                    "sample_values": samples,
                    "reason": "Header 'Contact' contained person name values",
                }
                assigned_canonicals.add("full_name")
                continue
            else:
                # Ambiguous without clear winner
                results[h] = {
                    "field": "",
                    "is_custom": False,
                    "confidence": 55,
                    "confidence_level": "LOW",
                    "sample_values": samples,
                    "reason": "Column 'Contact' is ambiguous — please select Phone, Email, or Name",
                }
                continue

        # 3. Direct alias match
        direct_match = NORMALIZED_ALIAS_MAP.get(compressed_h)
        if direct_match:
            # boost confidence if samples confirm it
            conf = 97
            if direct_match == "phone" and any(_is_probable_phone(s) for s in samples):
                conf = 99
            elif direct_match == "email" and any(_is_probable_email(s) for s in samples):
                conf = 99
            elif direct_match in ("full_name", "first_name", "last_name") and any(_is_probable_name(s) for s in samples):
                conf = 97

            results[h] = {
                "field": direct_match,
                "is_custom": False,
                "confidence": conf,
                "confidence_level": "HIGH",
                "sample_values": samples,
                "reason": f"Exact alias match for '{CANONICAL_FIELDS[direct_match]['label']}'",
            }
            assigned_canonicals.add(direct_match)
            continue

        # 4. Partial / Substring / Token matching
        tokens = set(norm_h.split())
        partial_hit = None
        partial_conf = 0

        # Phone check
        if any(t in ("phone", "mobile", "cell", "tel", "whatsapp", "call") for t in tokens):
            if any(_is_probable_phone(s) for s in samples):
                partial_hit = "phone"
                partial_conf = 92
            else:
                partial_hit = "phone"
                partial_conf = 78
        # Email check
        elif any(t in ("email", "mail") for t in tokens):
            partial_hit = "email"
            partial_conf = 95 if any(_is_probable_email(s) for s in samples) else 80
        # Name checks
        elif "first" in tokens and "name" in tokens:
            partial_hit = "first_name"
            partial_conf = 94
        elif "last" in tokens and "name" in tokens:
            partial_hit = "last_name"
            partial_conf = 94
        elif "name" in tokens:
            partial_hit = "full_name"
            partial_conf = 90
        # Company / Org
        elif any(t in ("company", "org", "corp", "business", "employer") for t in tokens):
            partial_hit = "company"
            partial_conf = 90
        # Job title
        elif any(t in ("title", "role", "designation", "position") for t in tokens):
            partial_hit = "job_title"
            partial_conf = 88
        # Location
        elif "city" in tokens:
            partial_hit = "city"
            partial_conf = 90
        elif "state" in tokens or "province" in tokens:
            partial_hit = "state"
            partial_conf = 90
        elif "country" in tokens:
            partial_hit = "country"
            partial_conf = 90
        elif "note" in tokens or "notes" in tokens or "comment" in tokens:
            partial_hit = "notes"
            partial_conf = 90

        if partial_hit and partial_hit not in assigned_canonicals:
            results[h] = {
                "field": partial_hit,
                "is_custom": False,
                "confidence": partial_conf,
                "confidence_level": "HIGH" if partial_conf >= 90 else "MEDIUM",
                "sample_values": samples,
                "reason": f"Pattern match for {CANONICAL_FIELDS[partial_hit]['label']}",
            }
            assigned_canonicals.add(partial_hit)
            continue

        # 5. Data-driven inference for remaining columns
        if any(_is_probable_email(s) for s in samples):
            results[h] = {
                "field": "email",
                "is_custom": False,
                "confidence": 85,
                "confidence_level": "MEDIUM",
                "sample_values": samples,
                "reason": "Values look like email addresses",
            }
            assigned_canonicals.add("email")
            continue

        if any(_is_probable_phone(s) for s in samples) and "phone" not in assigned_canonicals:
            results[h] = {
                "field": "phone",
                "is_custom": False,
                "confidence": 82,
                "confidence_level": "MEDIUM",
                "sample_values": samples,
                "reason": "Values look like phone numbers",
            }
            assigned_canonicals.add("phone")
            continue

        # 6. Unmapped column -> default to custom field with clean slug
        slug = re.sub(r"[^a-z0-9_]", "", norm_h.replace(" ", "_"))
        if not slug or slug[0].isdigit():
            slug = f"custom_{slug or 'field'}"
        results[h] = {
            "field": f"custom_fields.{slug}",
            "is_custom": True,
            "confidence": 65,
            "confidence_level": "LOW",
            "sample_values": samples,
            "reason": "Unrecognized column; preserved as custom field",
        }

    return results


def normalize_e164_phone(raw_phone: str, default_country: str = "US") -> tuple[str | None, str | None]:
    """Normalize phone number to E.164.
    Returns: (e164_string or None, error_message or None)
    """
    if not raw_phone or not str(raw_phone).strip():
        return None, "Empty phone number"

    cleaned = str(raw_phone).strip()

    try:
        import phonenumbers

        region = (default_country or "US").upper()
        # Parse using phonenumbers library
        parsed = phonenumbers.parse(cleaned, region)
        if not phonenumbers.is_possible_number(parsed):
            return None, f"Invalid phone length or format: '{raw_phone}'"
        if not phonenumbers.is_valid_number(parsed):
            # Sometimes local formats pass is_possible; check if valid in region
            pass
        formatted = phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164)
        return formatted, None
    except Exception:
        pass

    # Fallback normalization regex if phonenumbers fails or throws
    digits = re.sub(r"[^\d]", "", cleaned)
    if not digits:
        return None, f"No digits in phone number: '{raw_phone}'"

    country_code = DEFAULT_COUNTRY_CODES.get((default_country or "US").upper(), "1")

    if cleaned.startswith("+"):
        if 8 <= len(digits) <= 15:
            return f"+{digits}", None
        return None, f"Invalid international number length: '{raw_phone}'"

    # Digits without plus
    if len(digits) == 10:
        return f"+{country_code}{digits}", None
    elif len(digits) == 11 and digits.startswith(country_code):
        return f"+{digits}", None
    elif 8 <= len(digits) <= 15:
        # If it looks like it already includes country code
        return f"+{digits}", None

    return None, f"Malformed phone number: '{raw_phone}'"


def derive_names(contact_data: dict[str, Any]) -> dict[str, Any]:
    """Derive first/last name from full_name or full_name from first+last without destroying originals."""
    out = dict(contact_data)
    first = str(out.get("first_name") or "").strip()
    last = str(out.get("last_name") or "").strip()
    full = str(out.get("full_name") or "").strip()

    if full and not first and not last:
        parts = full.split(None, 1)
        out["first_name"] = parts[0]
        out["last_name"] = parts[1] if len(parts) > 1 else ""
    elif (first or last) and not full:
        out["full_name"] = f"{first} {last}".strip()

    return out


def normalize_and_validate_contacts(
    rows: list[dict[str, Any]],
    mapping: dict[str, str],
    default_country: str = "US",
    duplicate_strategy: str = "keep_first",
    source_file_name: str = "manual_import",
) -> dict[str, Any]:
    """Transform uploaded raw rows through the tenant's chosen mapping.
    Handles phone validation, deduplication, custom fields preservation, and raw data capture.
    """
    valid_contacts: list[dict[str, Any]] = []
    invalid_contacts: list[dict[str, Any]] = []
    duplicate_contacts: list[dict[str, Any]] = []

    # Map for deduplication: normalized_phone -> list of contact dicts
    phone_buckets: dict[str, list[dict[str, Any]]] = {}

    for row_idx, raw_row in enumerate(rows, start=1):
        contact_id = str(uuid.uuid4())
        canonical: dict[str, Any] = {
            "id": contact_id,
            "custom_fields": {},
            "raw_data": raw_row,
            "source": {"file": source_file_name, "row": row_idx},
        }

        # Apply mapping
        raw_phone_val = ""
        for src_col, target_field in mapping.items():
            if not target_field or target_field == "ignore":
                continue
            val = str(raw_row.get(src_col) or "").strip()

            if target_field == "phone":
                raw_phone_val = val
            elif target_field.startswith("custom_fields."):
                custom_key = target_field.split(".", 1)[1]
                canonical["custom_fields"][custom_key] = val
            elif target_field in CANONICAL_FIELDS:
                canonical[target_field] = val
            else:
                # Custom field without prefix
                canonical["custom_fields"][target_field] = val

        # Preserve unmapped columns in custom_fields so AI voice agent never loses CSV columns
        for src_col, val in raw_row.items():
            if src_col not in mapping or mapping[src_col] == "ignore":
                norm_col = re.sub(r"[^a-z0-9_]", "", src_col.lower().replace(" ", "_"))
                if norm_col and norm_col not in canonical["custom_fields"]:
                    canonical["custom_fields"][norm_col] = str(val).strip()

        # Derive names
        canonical = derive_names(canonical)

        # Phone normalization & validation
        if not raw_phone_val:
            invalid_contacts.append({
                "source_row": row_idx,
                "raw_data": raw_row,
                "reason": "Missing phone number in mapped column",
            })
            continue

        e164, err = normalize_e164_phone(raw_phone_val, default_country)
        if not e164:
            invalid_contacts.append({
                "source_row": row_idx,
                "raw_data": raw_row,
                "phone_attempted": raw_phone_val,
                "reason": err or "Invalid phone number",
            })
            continue

        canonical["phone"] = e164

        # Group by normalized phone for duplicate detection
        if e164 not in phone_buckets:
            phone_buckets[e164] = []
        phone_buckets[e164].append(canonical)

    # Apply duplicate strategy
    for phone_num, bucket in phone_buckets.items():
        if len(bucket) == 1:
            valid_contacts.append(bucket[0])
        else:
            # Duplicates found!
            if duplicate_strategy == "keep_last":
                chosen = bucket[-1]
                valid_contacts.append(chosen)
                for dup in bucket[:-1]:
                    duplicate_contacts.append({
                        "phone": phone_num,
                        "source_row": dup["source"]["row"],
                        "name": dup.get("full_name") or dup.get("first_name") or "",
                        "kept": False,
                    })
            elif duplicate_strategy == "merge":
                # Merge non-empty fields into the first contact
                merged = dict(bucket[0])
                merged_custom = dict(merged.get("custom_fields") or {})
                for later in bucket[1:]:
                    for k, v in later.items():
                        if v and not merged.get(k):
                            merged[k] = v
                    for ck, cv in (later.get("custom_fields") or {}).items():
                        if cv and not merged_custom.get(ck):
                            merged_custom[ck] = cv
                merged["custom_fields"] = merged_custom
                valid_contacts.append(merged)
                for dup in bucket[1:]:
                    duplicate_contacts.append({
                        "phone": phone_num,
                        "source_row": dup["source"]["row"],
                        "name": dup.get("full_name") or "",
                        "kept": "merged",
                    })
            elif duplicate_strategy == "remove":
                # Remove all copies
                for dup in bucket:
                    duplicate_contacts.append({
                        "phone": phone_num,
                        "source_row": dup["source"]["row"],
                        "name": dup.get("full_name") or "",
                        "kept": False,
                    })
            else:
                # Default: keep_first
                chosen = bucket[0]
                valid_contacts.append(chosen)
                for dup in bucket[1:]:
                    duplicate_contacts.append({
                        "phone": phone_num,
                        "source_row": dup["source"]["row"],
                        "name": dup.get("full_name") or dup.get("first_name") or "",
                        "kept": False,
                    })

    return {
        "total_rows": len(rows),
        "valid_count": len(valid_contacts),
        "invalid_count": len(invalid_contacts),
        "duplicate_count": len(duplicate_contacts),
        "valid_contacts": valid_contacts,
        "invalid_contacts": invalid_contacts,
        "duplicate_contacts": duplicate_contacts,
    }


def extract_template_variables_from_text(*texts: str) -> list[str]:
    """Find all {{contact.*}} and {{...}} placeholders in agent prompts."""
    variables: list[str] = []
    pattern = re.compile(r"\{\{\s*([a-zA-Z0-9_.]+)\s*\}\}")
    for t in texts:
        if not t:
            continue
        for m in pattern.finditer(t):
            var_name = m.group(1).strip()
            if var_name not in variables:
                variables.append(var_name)
    return variables


def resolve_contact_variables(contact: dict[str, Any]) -> dict[str, str]:
    """Resolve full runtime variable map for one contact."""
    out: dict[str, str] = {}
    full_name = str(contact.get("full_name") or "").strip()
    first_name = str(contact.get("first_name") or "").strip()
    last_name = str(contact.get("last_name") or "").strip()
    phone = str(contact.get("phone") or "").strip()
    email = str(contact.get("email") or "").strip()
    company = str(contact.get("company") or "").strip()
    job_title = str(contact.get("job_title") or "").strip()
    city = str(contact.get("city") or "").strip()
    state = str(contact.get("state") or "").strip()
    country = str(contact.get("country") or "").strip()
    notes = str(contact.get("notes") or "").strip()

    # Standard built-in variables
    out["contact.full_name"] = full_name
    out["contact.first_name"] = first_name
    out["contact.last_name"] = last_name
    out["contact.phone"] = phone
    out["contact.email"] = email
    out["contact.company"] = company
    out["contact.job_title"] = job_title
    out["contact.city"] = city
    out["contact.state"] = state
    out["contact.country"] = country
    out["contact.notes"] = notes

    # Shorthand aliases used in voice prompts
    out["caller_name"] = first_name or full_name
    out["callback_phone"] = phone
    out["first_name"] = first_name or full_name
    out["last_name"] = last_name
    out["full_name"] = full_name
    out["phone"] = phone
    out["email"] = email
    out["company"] = company

    # Custom fields
    custom = contact.get("custom_fields") or {}
    for k, v in custom.items():
        val_str = str(v).strip()
        out[f"contact.custom_fields.{k}"] = val_str
        out[f"custom_fields.{k}"] = val_str
        out[k] = val_str

    return out


def validate_agent_variables_against_contacts(
    agent_texts: list[str],
    contacts: list[dict[str, Any]],
) -> dict[str, Any]:
    """Check if all variables in the agent script exist in the contact list.
    Returns: {
        "valid": bool,
        "variables_used": list[str],
        "missing_breakdown": { var_name: missing_count }
    }
    """
    vars_used = extract_template_variables_from_text(*agent_texts)
    missing_breakdown: dict[str, int] = {}

    for var in vars_used:
        # Ignore system tags like {{business_name}}
        if var in ("business_name", "date", "time"):
            continue

        missing_count = 0
        for c in contacts:
            resolved = resolve_contact_variables(c)
            val = resolved.get(var)
            if not val:
                missing_count += 1

        if missing_count > 0:
            missing_breakdown[var] = missing_count

    return {
        "valid": len(missing_breakdown) == 0,
        "variables_used": vars_used,
        "missing_breakdown": missing_breakdown,
        "total_contacts": len(contacts),
    }


def render_agent_prompt_for_contact(prompt: str, contact: dict[str, Any]) -> str:
    """Render agent script replacing {{var}} tags with contact-specific values."""
    if not prompt:
        return ""
    resolved = resolve_contact_variables(contact)

    def replacer(match):
        key = match.group(1).strip()
        return resolved.get(key, match.group(0))

    return re.sub(r"\{\{\s*([a-zA-Z0-9_.]+)\s*\}\}", replacer, prompt)

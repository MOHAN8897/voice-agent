"""Workspace analytics and exports: tenant isolation, filters, formula escaping, xlsx."""
import csv
import io
import uuid
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import JSON
from sqlalchemy.dialects.postgresql import ARRAY, JSONB
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from server.db.models.entities import Agent, Call, CallAttempt, CallCallback, Tenant
from server.db.models.followup_models import FollowUpJob, LeadAssessment
from server.db.models.language_switch_models import LanguageSwitchRequest
from server.db.models.phase5_models import DncEntry
from server.db.models.saas_models import BillingWallet, BillingWalletTransaction, Lead
from server.services.saas import workspace_exports as exports
from server.services.saas.workspace_analytics import outcome_bucket, resolve_range, workspace_analytics

NOW = datetime.now(timezone.utc).replace(microsecond=0) - timedelta(minutes=5)
TABLES = [Tenant, Agent, Call, CallAttempt, CallCallback, Lead, LeadAssessment, FollowUpJob,
          LanguageSwitchRequest, DncEntry, BillingWallet, BillingWalletTransaction]


def _call(tid, aid, **kw):
    base = dict(call_id=uuid.uuid4(), tenant_id=tid, agent_id=aid, combination_id="c", storage_path="t",
                channel="pstn", direction="inbound", status="answered", started_at=NOW,
                ended_at=NOW + timedelta(seconds=120), duration_sec=120, is_test=False, status_tags=[])
    base.update(kw)
    return Call(**base)


async def _seed_tenant(session, label):
    tid, aid = uuid.uuid4(), uuid.uuid4()
    session.add_all([Tenant(tenant_id=tid, name=label), Agent(agent_id=aid, tenant_id=tid, name=f"{label} agent")])
    await session.flush()
    answered = _call(tid, aid, status_tags=["intent:pricing", "custom:solar_panels"])
    session.add_all([
        answered,
        _call(tid, aid, direction="outbound", status="outbound", status_tags=["intent:pricing"]),
        _call(tid, aid, direction="outbound", status="declined", end_reason="busy", duration_sec=0),
        _call(tid, aid, status="answered", is_test=True),
        CallAttempt(attempt_id=uuid.uuid4(), tenant_id=tid, agent_id=aid, provider_call_control_id=f"{label}-1",
                    direction="outbound", status="missed", from_number="+918000000000", to_number="+919398247132",
                    started_at=NOW, ended_at=NOW),
        CallAttempt(attempt_id=uuid.uuid4(), tenant_id=tid, agent_id=aid, provider_call_control_id=f"{label}-2",
                    direction="inbound", status="voicemail", from_number="+919398247133", to_number="+918000000000",
                    started_at=NOW, ended_at=NOW),
    ])
    lead = Lead(lead_id=uuid.uuid4(), tenant_id=tid, agent_id=str(aid), name="=HYPERLINK(\"http://x\")",
                phone="+919398247132", stage="qualified", created_at=NOW, updated_at=NOW)
    plain = Lead(lead_id=uuid.uuid4(), tenant_id=tid, agent_id=None, name=f"{label} walk-in", phone="9398247134",
                 stage="new", created_at=NOW, updated_at=NOW)
    session.add_all([lead, plain])
    await session.flush()
    session.add(LeadAssessment(tenant_id=tid, lead_id=lead.lead_id, source_call_id=str(answered.call_id),
                               assessment={}, updated_at=NOW))
    session.add_all([
        CallCallback(tenant_id=tid, agent_id=str(aid), to_e164="+919398247132", dial_request_id=f"{label}-cb1",
                     mode="requested", status="requested", created_at=NOW),
        CallCallback(tenant_id=tid, agent_id=str(aid), to_e164="+919398247135", dial_request_id=f"{label}-cb2",
                     mode="requested", status="completed", created_at=NOW),
        LanguageSwitchRequest(tenant_id=tid, agent_id=str(aid), phone_e164="+919398247132", requested_language="hi",
                              original_language="te", status="pending", created_at=NOW),
        LanguageSwitchRequest(tenant_id=tid, agent_id=str(aid), phone_e164="+919398247136", requested_language="ta",
                              original_language="te", status="done", created_at=NOW),
        FollowUpJob(tenant_id=tid, agent_id=aid, origin_key=f"{label}-f1", source_id="x", kind="callback",
                    phone="+919398247132", status="scheduled", reason="r", payload={}, expires_at=NOW + timedelta(days=3),
                    created_at=NOW, updated_at=NOW),
        FollowUpJob(tenant_id=tid, agent_id=aid, origin_key=f"{label}-f2", source_id="y", kind="recall",
                    phone="+919398247137", status="completed", reason="call_connected", payload={},
                    expires_at=NOW + timedelta(days=3), created_at=NOW, updated_at=NOW),
        FollowUpJob(tenant_id=tid, agent_id=aid, origin_key=f"{label}-f3", source_id="z", kind="recall",
                    phone="+919398247138", status="expired", reason="request_expired", payload={},
                    expires_at=NOW, created_at=NOW, updated_at=NOW),
        DncEntry(tenant_id=tid, phone_e164="+919398247139", source="live_call_opt_out", active=True, added_at=NOW),
        BillingWallet(tenant_id=tid, balance_cents=0, currency="inr"),
        BillingWalletTransaction(tenant_id=tid, amount_cents=0, amount_inr_paise=-1250, kind="call_usage",
                                 reference_id=f"call:{answered.call_id}", created_at=NOW),
        BillingWalletTransaction(tenant_id=tid, amount_cents=0, amount_inr_paise=50000, kind="topup",
                                 reference_id=f"topup:{label}", created_at=NOW),
    ])
    return SimpleNamespace(tid=str(tid), aid=str(aid))


@pytest.fixture
async def db(monkeypatch, tmp_path):
    from server.db import connection
    from server.services.saas import agent_performance, workspace_analytics

    changed = [(c, c.type) for m in TABLES for c in m.__table__.columns if isinstance(c.type, (ARRAY, JSONB))]
    for c, _ in changed:
        c.type = JSON()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'ws.db'}")
    try:
        async with engine.begin() as conn:
            for model in TABLES:
                await conn.run_sync(lambda sync, table=model.__table__: table.create(sync))
        factory = async_sessionmaker(engine, expire_on_commit=False)
        for module in (connection, agent_performance, workspace_analytics, exports):
            monkeypatch.setattr(module, "get_session_factory", lambda: factory)
        async with factory() as session:
            a = await _seed_tenant(session, "A")
            b = await _seed_tenant(session, "B")
            await session.commit()
        yield SimpleNamespace(a=a, b=b)
    finally:
        for c, original in changed:
            c.type = original
        await engine.dispose()


async def test_analytics_counts_only_the_callers_tenant(db):
    day = NOW.astimezone(timezone(timedelta(hours=5, minutes=30))).date().isoformat()
    data = await workspace_analytics(db.a.tid, range_key="custom", since=day, until=day)
    calls = data["calls"]
    # 3 real calls + 2 unanswered attempts; the practice call is excluded.
    assert calls["total"] == 5
    assert calls["outcomes"]["answered"] == 2
    assert calls["outcomes"]["busy"] == 1
    assert calls["outcomes"]["no_answer"] == 1
    assert calls["outcomes"]["voicemail"] == 1
    assert sum(d["calls"] for d in calls["perDay"]) == 5
    assert data["leads"]["total"] == 2
    assert data["leads"]["bySource"] == {"inbound_call": 1, "manual": 1}
    assert data["callbacks"] == {**data["callbacks"], "requested": 2, "completed": 1, "pending": 1}
    assert data["languageCallbacks"]["requested"] == 2 and data["languageCallbacks"]["completed"] == 1
    assert {row["language"] for row in data["languageCallbacks"]["byLanguage"]} == {"hi", "ta"}
    assert (data["followups"]["scheduled"], data["followups"]["done"], data["followups"]["failed"]) == (1, 1, 1)
    assert data["dnc"]["active"] == 1 and data["dnc"]["added"] == 1
    assert data["spend"]["currency"] == "INR" and data["spend"]["total"] == 12.5
    assert data["spend"]["byAgent"][0]["name"] == "A agent"
    assert data["topTags"][0] == {"tag": "intent:pricing", "count": 2}


async def test_analytics_rejects_another_tenants_agent(db):
    with pytest.raises(HTTPException) as err:
        await workspace_analytics(db.a.tid, agent_id=db.b.aid, range_key="7d")
    assert err.value.status_code == 404


def test_range_rules():
    start, end, days = resolve_range("custom", "2026-10-01", "2026-10-03")
    assert len(days) == 3 and (end - start) == timedelta(days=3)
    assert start.isoformat() == "2026-09-30T18:30:00+00:00"  # IST midnight
    with pytest.raises(HTTPException):
        resolve_range("custom", "2026-10-05", "2026-10-01")
    with pytest.raises(HTTPException):
        resolve_range("custom", "2024-01-01", "2026-01-01")


def test_outcome_buckets():
    assert outcome_bucket("declined", "outbound", "busy") == "busy"
    assert outcome_bucket("missed", "inbound") == "no_answer"
    assert outcome_bucket("failed", "outbound", "unreachable") == "unreachable"
    assert outcome_bucket("outbound", "outbound") == "answered"


async def test_calls_export_filters_and_isolation(db):
    data = await exports.collect("calls", db.a.tid, exports.Filters(direction="outbound"))
    assert data["total"] == 3  # two calls + the unanswered attempt, tenant A only
    tagged = await exports.collect("calls", db.a.tid, exports.Filters(tags="intent:pricing,custom:solar_panels", tagMatch="all"))
    assert tagged["total"] == 1
    anytag = await exports.collect("calls", db.a.tid, exports.Filters(tags="intent:pricing,custom:solar_panels"))
    assert anytag["total"] == 2
    busy = await exports.collect("calls", db.a.tid, exports.Filters(outcome="busy"))
    assert busy["total"] == 1 and busy["rows"][0][5] == "Busy"
    missed = await exports.collect("calls", db.a.tid, exports.Filters(outcome="no_answer"))
    assert missed["rows"][0][3] == "+91 93982 47132"  # formatted contact, not our caller ID


async def test_leads_dnc_followup_wallet_filters(db):
    qualified = await exports.collect("leads", db.a.tid, exports.Filters(stage="qualified"))
    assert qualified["total"] == 1 and qualified["rows"][0][2] == "+91 93982 47132"
    by_tag = await exports.collect("leads", db.a.tid, exports.Filters(tags="custom:solar_panels"))
    assert by_tag["total"] == 1 and "Solar panels" in by_tag["rows"][0][7]
    assert (await exports.collect("dnc", db.a.tid, exports.Filters(source="manual")))["total"] == 0
    assert (await exports.collect("dnc", db.a.tid, exports.Filters(source="live_call_opt_out")))["total"] == 1
    assert (await exports.collect("followups", db.a.tid, exports.Filters(status="failed")))["total"] == 1
    assert (await exports.collect("followups", db.a.tid, exports.Filters(kind="recall")))["total"] == 2
    assert (await exports.collect("callbacks", db.a.tid, exports.Filters(status="requested")))["total"] == 1
    assert (await exports.collect("language_callbacks", db.a.tid, exports.Filters(language="hi")))["total"] == 1
    wallet = await exports.collect("wallet", db.a.tid, exports.Filters(kind="call_usage"))
    assert wallet["rows"][0][2] == -12.5 and wallet["rows"][0][4] == "A agent"
    with pytest.raises(HTTPException):
        await exports.collect("leads", db.a.tid, exports.Filters(agentId=db.b.aid))


def test_formula_cells_are_neutralised():
    for bad in ("=1+1", "+cmd", "-2+3", "@SUM(A1)", "\t=x"):
        assert exports.safe_cell(bad).startswith("'")
    assert exports.safe_cell("Ravi") == "Ravi"
    assert exports.safe_cell(-12.5) == -12.5


async def test_csv_and_xlsx_outputs_are_safe_and_valid(db):
    data = await exports.collect("leads", db.a.tid, exports.Filters(stage="qualified"))
    text = "".join(exports.csv_chunks(data["columns"], data["rows"])).lstrip("﻿")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0][0] == "Created (IST)"
    assert rows[1][1].startswith("'=HYPERLINK")
    book = load_workbook(io.BytesIO(exports.xlsx_bytes(data["columns"], data["rows"], "Leads")))
    sheet = book.active
    assert sheet.title == "Leads"
    # Excel stores the text as a string cell, so it is shown as typed and never run;
    # no quote prefix is needed there (phone numbers read "+91 …", not "'+91 …").
    assert sheet.cell(row=2, column=2).value.startswith("=HYPERLINK")
    assert sheet.cell(row=2, column=2).data_type == "s"
    assert sheet.max_row == 2


def test_wallet_export_needs_billing_permission():
    from server.routes.app_analytics import _tenant

    viewer = SimpleNamespace(role="customer_viewer", tenant_id=uuid.uuid4(), email="v@example.com")
    assert _tenant(viewer, "wallet") == viewer.tenant_id
    outsider = SimpleNamespace(role="nobody", tenant_id=uuid.uuid4(), email="n@example.com")
    with pytest.raises(HTTPException) as err:
        _tenant(outsider, "calls")
    assert err.value.status_code == 403
    with pytest.raises(HTTPException):
        _tenant(viewer, "secrets")
